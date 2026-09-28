"""Leitura de fichas técnicas no modelo de planilha "FICHA TÉCNICA
OPERACIONAL" (uma ficha por aba), usado pela cozinha:

    E4  Produto:<nome>            A9  Ingredientes | E9 Peso Bruto | G9 Peso Líquido
    E5  Rendimento:<g>            A10.. um ingrediente por linha, até TOTAL
    E6  Número de porções:<n>     MODO DE PREPARO  (texto quebrado em várias linhas)
    E7  Porção em gramas:<g>      OBSERVAÇÕES      (Equipamentos, Tempo de cocção,
                                                    Outras orientações, Refrigeração)

As posições variam um pouco de aba para aba (linhas a mais ou a menos), então
a leitura procura pelos rótulos em vez de usar linhas fixas.

Só lê e normaliza texto — não grava nada. Quem decide nome final, categoria
e de-para de insumos é o chamador (ver scripts/importar_fichas_brava.py).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class ItemPlanilha:
    nome: str
    peso_bruto: float | None
    peso_liquido: float | None


@dataclass
class FichaPlanilha:
    arquivo: str
    aba: str
    produto: str
    rendimento_g: float | None
    rendimento_texto: str
    porcoes: str
    porcao_g: str
    itens: list[ItemPlanilha] = field(default_factory=list)
    passos: list[str] = field(default_factory=list)
    avisos_preparo: list[str] = field(default_factory=list)  # linhas "OBS:" dentro do preparo
    equipamentos: str = ""
    tempo_coccao: str = ""
    orientacoes: str = ""
    refrigeracao: str = ""

    @property
    def vazia(self) -> bool:
        """Aba ainda não preenchida: sem ingrediente com peso."""
        return not any(i.peso_bruto for i in self.itens)


def sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def chave(texto: str) -> str:
    """Forma comparável de um nome: sem acento, maiúsculo e espaços simples."""
    return " ".join(sem_acento(texto).upper().split())


def frase(texto: str) -> str:
    """Texto em CAIXA ALTA -> frase normal ("REFOGAR O ALHO,A CEBOLA" ->
    "Refogar o alho, a cebola"). Ajusta o espaço depois de vírgula, que nas
    fichas quase sempre vem colado."""
    texto = " ".join(texto.split()).lower()
    texto = re.sub(r"\s*,\s*", ", ", texto)
    texto = re.sub(r"(?<=[a-zà-ú)])\s*:\s*(?=\S)", ": ", texto)
    texto = re.sub(r"(?<=[a-zà-ú])\s*/\s*(?=[a-zà-ú])", " / ", texto)
    texto = re.sub(r"(\d)\s*°\s*c\b", r"\1°C", texto)
    return texto[:1].upper() + texto[1:] if texto else texto


def numero_br(texto: str) -> float | None:
    """ "425" -> 425; "3.606" (milhar com ponto) -> 3606; "0,5" -> 0.5."""
    texto = texto.strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
        return float(texto.replace(".", ""))
    try:
        return float(texto.replace(",", "."))
    except ValueError:
        return None


def _valor_rotulado(celula: object, rotulo: str) -> str | None:
    """ "Produto:MOLHO BECHAMEL" com rotulo "Produto" -> "MOLHO BECHAMEL"."""
    if isinstance(celula, str) and chave(celula).startswith(chave(rotulo)):
        return celula.split(":", 1)[1].strip() if ":" in celula else ""
    return None


def _num(celula: object) -> float | None:
    if isinstance(celula, (int, float)):
        return float(celula)
    if isinstance(celula, str) and celula.strip():
        return numero_br(celula)
    return None


def _passos(linhas: list[str]) -> tuple[list[str], list[str]]:
    """Junta as linhas do modo de preparo (quebradas no meio da frase pela
    largura da célula) e reparte por frase. A numeração "1 -" das fichas é por
    linha, não por passo, então é descartada."""
    avisos, corpo = [], []
    for linha in linhas:
        linha = re.sub(r"^\s*\d+\s*-\s*(?!\d)", "", linha).strip()  # "1 -" sim, "40-50 MIN" não
        if not linha:
            continue
        if chave(linha).startswith("OBS"):
            avisos.append(frase(linha.split(":", 1)[-1]))
        else:
            corpo.append(linha)
    texto = " ".join(corpo)
    # Quebra em frase (ponto, mesmo colado: "GORDURA.DEGLASAR") e antes de
    # rótulos de etapa ("FORMAR O ROUX:", "DEGLASAGEM:").
    texto = re.sub(r"\.\s*(?=[A-ZÀ-Ú])", ".\n", texto)
    texto = re.sub(r"\s+((?:[A-ZÀ-Ú]+ ){0,2}[A-ZÀ-Ú]{4,}:)", r"\n\1", texto)
    passos = [frase(p) for p in texto.split("\n") if p.strip(" .")]
    return passos, avisos


def ler_aba(ws: Worksheet, arquivo: str) -> FichaPlanilha:
    celulas = {(c.row, c.column): c.value for linha in ws.iter_rows() for c in linha if c.value not in (None, "")}

    def col_a(linha: int) -> str:
        valor = celulas.get((linha, 1))
        return valor.strip() if isinstance(valor, str) else ""

    cab: dict[str, str] = {}
    for valor in celulas.values():
        for rotulo in ("Produto", "Rendimento", "Número de porções", "Porção em gramas"):
            achado = _valor_rotulado(valor, rotulo)
            if achado is not None:
                cab[rotulo] = achado

    ultima = ws.max_row
    inicio = next(r for r in range(1, ultima + 1) if chave(col_a(r)).startswith("INGREDIENTES"))
    preparo = next(r for r in range(inicio, ultima + 1) if chave(col_a(r)) == "MODO DE PREPARO")
    observacoes = next((r for r in range(preparo, ultima + 1) if chave(col_a(r)) == "OBSERVACOES"), ultima + 1)

    itens = []
    for r in range(inicio + 1, preparo):
        nome = col_a(r)
        if not nome or chave(nome) == "TOTAL":
            continue
        itens.append(ItemPlanilha(" ".join(nome.split()), _num(celulas.get((r, 5))), _num(celulas.get((r, 7)))))

    passos, avisos = _passos([col_a(r) for r in range(preparo + 1, observacoes)])

    extras = {"equipamentos": "", "tempo_coccao": "", "orientacoes": "", "refrigeracao": ""}
    rotulos = {
        "EQUIPAMENTOS": "equipamentos",
        "TEMPO DE COCCAO": "tempo_coccao",
        "OUTRAS ORIENTACOES": "orientacoes",
        "REFRIGERACAO": "refrigeracao",
    }
    for r in range(observacoes + 1, ultima + 1):
        linha = col_a(r)
        for prefixo, campo in rotulos.items():
            if chave(linha).startswith(prefixo):
                extras[campo] = linha.split(":", 1)[1].strip() if ":" in linha else ""

    rend_texto = cab.get("Rendimento", "")
    return FichaPlanilha(
        arquivo=arquivo,
        aba=ws.title,
        produto=cab.get("Produto", ws.title).strip(),
        rendimento_g=numero_br(rend_texto) if rend_texto else None,
        rendimento_texto=rend_texto,
        porcoes=cab.get("Número de porções", ""),
        porcao_g=cab.get("Porção em gramas", ""),
        itens=itens,
        passos=passos,
        avisos_preparo=avisos,
        **extras,
    )


def ler_planilha(caminho: Path) -> list[FichaPlanilha]:
    wb = load_workbook(caminho, data_only=True)
    return [ler_aba(ws, caminho.name) for ws in wb.worksheets]


def dias_refrigeracao(texto: str) -> int | None:
    """ "48H 2 - 4°C" -> 2; "5 DIAS" -> 5. None se não der pra entender."""
    t = chave(texto)
    if m := re.search(r"(\d+)\s*H\b", t):
        return max(1, int(m.group(1)) // 24)
    if m := re.search(r"(\d+)\s*DIAS?\b", t):
        return int(m.group(1))
    return None
